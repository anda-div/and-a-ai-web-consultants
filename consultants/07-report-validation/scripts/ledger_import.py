# -*- coding: utf-8 -*-
"""過去資料の取り込み（汎用）

毎月の改善提案からは、次の3つを外さなければならない。

    1. 過去に実施した提案      効果の良し悪しに関係なく外す
    2. いま実装中・実装予定の提案
    3. 一度検討されて却下、またはうやむやになっている提案

外さずに出すと、**「過去の経緯を確認していない支援会社」** に見える。
品質ではなく信頼の問題で、一度その印象がつくとレポート全体の読まれ方が変わる。

重複を避ける判定機構（`--similar` `--angle` `--target`）は既にある。
足りないのは**入口**、つまりクライアントの手元にある資料を台帳へ入れる経路だった。

機械にできることと、できないこと
--------------------------------
PowerPointやPDFから「これは提案である」を機械が判定することはできない。
無理に判定させると、台帳に意味のない行が積まれる。**空の台帳より悪い。**

そこで工程を2つに分ける。

    1. ledger.py --import      機械が集める
                               ・どのファイルがどの箱に入っているか（＝申告された状態）
                               ・中の文字をすべて取り出す
                               ・URL・切り口・日付の候補を拾っておく
                               ・取り込み済みのものは飛ばす
    2. ledger.py --import-apply  AIが読んで埋めた下書きを、台帳へ入れる
                               ・target × angle で重複を判定する
                               ・状態ごとの必須項目を検査する
                               ・原本を _archive/ へ退避し、二度と読まない

真ん中（資料を読んで提案の形にする）はAIの仕事で、最後の承認は人がする。
`PROPOSAL_LEDGER.md` のF章に運用の全体がある。
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import re
import shutil
from datetime import date

# 抽出した文字のうち、下書きに載せる長さ。全文は別ファイルに書き出す
PREVIEW_CHARS = 400
EXTRACT_DIR = "_extract"


# ---------------------------------------------------------------- 文字の取り出し
def _read_text(path: str) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp932"):
        try:
            with io.open(path, encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with io.open(path, encoding="utf-8", errors="replace") as f:
        return f.read()


def _read_pptx(path: str) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for n, slide in enumerate(prs.slides, 1):
        out.append(f"--- スライド {n} ---")

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type is not None and sh._element.tag.endswith("}grpSp"):
                    walk(sh.shapes)
                    continue
                if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                    out.append(sh.text_frame.text.strip())
                if getattr(sh, "has_table", False):
                    for row in sh.table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        if any(cells):
                            out.append(" | ".join(cells))

        walk(slide.shapes)
        # ノートに経緯が書かれていることがある
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                out.append(f"［ノート］{note}")
    return "\n".join(out)


def _read_docx(path: str) -> str:
    import docx
    d = docx.Document(path)
    out = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    return "\n".join(out)


def _read_xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"--- シート {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if cells:
                out.append(" | ".join(cells))
    wb.close()
    return "\n".join(out)


def _read_pdf(path: str) -> str:
    from pypdf import PdfReader
    r = PdfReader(path)
    out = []
    for n, page in enumerate(r.pages, 1):
        t = (page.extract_text() or "").strip()
        if t:
            out.append(f"--- ページ {n} ---")
            out.append(t)
    return "\n".join(out)


READERS = {".md": _read_text, ".txt": _read_text, ".pptx": _read_pptx,
           ".docx": _read_docx, ".xlsx": _read_xlsx, ".pdf": _read_pdf}


def extract(path: str) -> tuple[str, str]:
    """(取り出した文字, 読めなかった理由) を返す。

    読めないことを黙って0文字として扱わない。**「読めなかった」を
    「中身が無かった」にしない。** 下書きに理由を残し、人が見に行けるようにする。
    """
    ext = os.path.splitext(path)[1].lower()
    fn = READERS.get(ext)
    if fn is None:
        return "", f"{ext or '拡張子なし'} は読めません"
    try:
        return fn(path), ""
    except ImportError as e:
        return "", f"読み取りに必要なライブラリがありません（{e.name}）"
    except Exception as e:
        return "", f"読めませんでした（{type(e).__name__}: {e}）"


# ---------------------------------------------------------------- 手がかり
_URL = re.compile(r"https?://[^\s　「」（）()<>\"']+")
_PATH = re.compile(r"(?<![\w./])/[A-Za-z0-9][\w\-./]{1,60}")
_DATE = re.compile(r"(20\d{2})\s*[-/年]\s*(\d{1,2})\s*[-/月]?")


def hints(text: str, catalog: dict | None) -> dict:
    """target・angle・実施時期の候補を拾う。**確定はしない。**

    機械に決めさせると外れたときに気づけない。候補として並べ、
    AIが本文を読んで選び、人が確認シートで直す。
    """
    targets = []
    for m in list(_URL.finditer(text))[:40] + list(_PATH.finditer(text))[:40]:
        v = m.group(0).rstrip("。、）)」.,")
        if v not in targets:
            targets.append(v)

    angles = []
    for a in ((catalog or {}).get("angles") or []):
        key, name = a.get("key", ""), a.get("name", "")
        if (key and key in text) or (name and name in text):
            if key not in angles:
                angles.append(key)

    dates = []
    for m in _DATE.finditer(text):
        y, mo = int(m.group(1)), int(m.group(2))
        if 1 <= mo <= 12:
            v = f"{y:04d}-{mo:02d}"
            if v not in dates:
                dates.append(v)

    return {"targets": targets[:12], "angles": angles[:8], "dates": sorted(dates)[:12]}


# ---------------------------------------------------------------- 取り込み口
def folder_specs(spec: dict) -> dict:
    """箱の定義だけを取り出す。

    JSONには説明のための `_comment`（文字列）が混ざる。仕様として舐めると
    文字列に .get を呼んで落ちる。**アンダースコア始まりと非辞書は外す。**
    """
    return {k: v for k, v in (spec.get("folders") or {}).items()
            if isinstance(v, dict) and not k.startswith("_")}


def sha1(path: str) -> str:
    h = hashlib.sha1()
    with io.open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def init_inbox(root: str, spec: dict, readme: str = "") -> list[str]:
    """取り込み口のフォルダを作る。既にあるものは触らない。"""
    made = []
    base = os.path.join(root, spec.get("dir_name", "_proposals_inbox"))
    names = list(folder_specs(spec))
    for name in sorted(names) + [spec.get("archive_dir", "_archive")]:
        d = os.path.join(base, name)
        if not os.path.isdir(d):
            os.makedirs(d)
            made.append(d)
        keep = os.path.join(d, ".gitkeep")
        if not os.path.exists(keep):
            io.open(keep, "w", encoding="utf-8").close()
    rp = os.path.join(base, "README.md")
    if readme and not os.path.exists(rp):
        with io.open(rp, "w", encoding="utf-8", newline="\n") as f:
            f.write(readme)
        made.append(rp)
    return made


def scan(root: str, spec: dict, done: set[str]) -> tuple[list[dict], list[dict]]:
    """取り込み口を見て、(未取り込みのファイル, 飛ばしたもの) を返す。

    done は取り込み済みの sha1。**同じ原本を二度読まない。**
    毎月PDFを読み直すと遅いうえ、読むたびに解釈がぶれる。
    """
    base = os.path.join(root, spec.get("dir_name", "_proposals_inbox"))
    folders = folder_specs(spec)
    ignore = set(spec.get("ignore_names") or [])
    archive = spec.get("archive_dir", "_archive")
    found, skipped = [], []
    if not os.path.isdir(base):
        return found, skipped
    for folder in sorted(folders):
        d = os.path.join(base, folder)
        if not os.path.isdir(d):
            continue
        for dirpath, dirnames, filenames in os.walk(d):
            dirnames[:] = [x for x in dirnames if x != archive]
            for name in sorted(filenames):
                if name in ignore or name.startswith("."):
                    continue
                p = os.path.join(dirpath, name)
                rel = os.path.relpath(p, base).replace("\\", "/")
                digest = sha1(p)
                rec = {"file": rel, "path": p, "sha1": digest, "folder": folder,
                       "status": folders[folder].get("status", "提案中"),
                       "needs_review": bool(folders[folder].get("needs_review")),
                       "require": folders[folder].get("require") or []}
                if digest in done:
                    skipped.append(rec)
                else:
                    found.append(rec)
    return found, skipped


def build_draft(root: str, spec: dict, catalog: dict | None,
                records: list[dict], out_dir: str) -> dict:
    """未取り込みのファイルから、AIが埋める下書きを作る。

    全文は `_extract/` に書き出し、下書きには先頭だけを載せる。
    50ページのPDFを下書きに流し込むと、肝心の手順が埋もれる。
    """
    ex_dir = os.path.join(out_dir, EXTRACT_DIR)
    os.makedirs(ex_dir, exist_ok=True)
    sources = []
    for r in records:
        text, why = extract(r["path"])
        stem = re.sub(r"[^\w.-]", "_", r["file"])[:80]
        ex_path = os.path.join(ex_dir, f"{r['sha1'][:8]}_{stem}.txt")
        with io.open(ex_path, "w", encoding="utf-8", newline="\n") as f:
            f.write(text)
        sources.append({
            "file": r["file"], "sha1": r["sha1"], "folder": r["folder"],
            "申告された状態": r["status"],
            "確認が要る": r["needs_review"],
            "必須項目": r["require"],
            "全文": os.path.relpath(ex_path, root).replace("\\", "/"),
            "文字数": len(text),
            "読めなかった理由": why,
            "冒頭": text[:PREVIEW_CHARS],
            "手がかり": hints(text, catalog),
        })
    return {
        "_手順": [
            "1. 各 sources の「全文」のファイルを読む。",
            "2. 提案を1件ずつ proposals に書く。**資料1本から複数件出てよい。**",
            "3. source には、その提案の出どころの file をそのまま書く。",
            "4. status は「申告された状態」を既定にする。本文を読んで違うと分かれば直す。",
            "5. target（対象URL）と angle（切り口カタログのキー）は必ず埋める。"
            "この2つが重複判定の軸で、空だと同じ提案を来月もう一度出す。",
            "6. 埋められないものは needs_ask に理由を書く。推測で埋めない。",
            "7. 実装済みは implemented_on を入れる。**「2025-10」のような概算でよい。**",
            "8. 保留は blocked_by（何待ちか）と revisit_on（いつ見直すか）を必ず入れる。"
            "「次期リニューアル時に」は却下ではなく保留＋期日。",
            "9. 読んだが提案が入っていなかった原本は、「提案なし」にファイル名を書く。"
            "書かないと入口に残り、次回もう一度出てくる（**黙って消さないため**）。",
            "10. 書き終えたら ledger.py --import-apply <このファイル> を実行する。",
        ],
        "提案なし": [],
        "_雛形": {
            "source": "", "title": "", "target": "", "angle": "",
            "status": "", "effort": "", "vendor": "", "metric": "",
            "implemented_on": "", "blocked_by": "", "revisit_on": "",
            "status_note": "", "created": "", "needs_ask": "",
        },
        "created": date.today().isoformat(),
        "sources": sources,
        "proposals": [],
    }


# ---------------------------------------------------------------- 台帳へ入れる
def check_row(row: dict, spec: dict) -> list[str]:
    """1件ぶんの検査。埋まっていないものを挙げる（**推測で埋めない**）。"""
    req = (spec.get("required_fields") or {}).get("always") or []
    bad = [k for k in req if not str(row.get(k) or "").strip()]
    extra = []
    for f in folder_specs(spec).values():
        if f.get("status") == row.get("status"):
            extra = f.get("require") or []
            break
    if row.get("status") == "保留":
        extra = list(set(extra) | {"blocked_by", "revisit_on"})
    if row.get("status") == "実装済み":
        extra = list(set(extra) | {"implemented_on"})
    bad += [k for k in extra if not str(row.get(k) or "").strip()]
    return sorted(set(bad))


def apply_draft(lg, draft: dict, spec: dict, *, period: str,
                archive_root: str = "", dry_run: bool = False) -> dict:
    """埋まった下書きを台帳へ入れる。

    重複は **target × angle** で見る。タイトルの類似では、言い換えられた
    同じ提案を取りこぼす。「料金を分かりやすく表示する」と「料金・最低保証を
    事前提示し、依頼フローを図解する」は、共通する文字がほとんど無い。
    """
    added, dup, bad = [], [], []
    used_sources = set()
    for row in draft.get("proposals") or []:
        missing = check_row(row, spec)
        if missing:
            bad.append((row, missing))
            continue
        hit = duplicate_of(lg, row.get("target", ""), row.get("angle", ""))
        if hit is not None:
            dup.append((row, hit))
            continue
        item = lg.add(
            period=row.get("created") or period,
            title=row["title"], target=row["target"], angle=row["angle"],
            status=row["status"], effort=row.get("effort", ""),
            vendor=row.get("vendor", ""), metric=row.get("metric", ""),
            status_note=row.get("status_note", ""),
            blocked_by=row.get("blocked_by", ""),
            revisit_on=row.get("revisit_on", ""),
            implemented_on=row.get("implemented_on") or None,
        )
        # どの原本から来たかを残す。後から「これはどこ発か」を追えるようにする
        item["source"] = row.get("source", "")
        item["imported_on"] = date.today().isoformat()
        added.append(item)
        if row.get("source"):
            used_sources.add(row["source"])

    # 片付いた原本だけを退避する。
    #
    # **埋まらなかった原本を退避してはいけない。** 台帳に入っていないのに
    # 入口から消え、sha1 も登録されるので --import が二度と拾わない。
    # 情報が黙って消える。片付いたとみなすのは次の3つだけ。
    #
    #   ・台帳に入った提案がある
    #   ・すでに台帳にある（重複）と分かった ── 中身は失われていない
    #   ・読んだうえで「提案は無い」と下書きに書いた（"提案なし"）
    #
    # 欠落のある原本は入口に残す。--import と --check が次回また言う。
    done_ok, blocked = set(), set()
    for row, _miss in bad:
        if row.get("source"):
            blocked.add(row["source"])
    for i in added:
        if i.get("source"):
            done_ok.add(i["source"])
    for row, _hit in dup:
        if row.get("source"):
            done_ok.add(row["source"])
    for f in (draft.get("提案なし") or []):
        done_ok.add(f)
    done_ok -= blocked

    moved, kept = [], []
    if not dry_run:
        reg = lg.data.setdefault("imports", [])
        known = {r.get("sha1") for r in reg}
        for s in draft.get("sources") or []:
            f = s.get("file", "")
            if f not in done_ok:
                kept.append(f)
                continue
            if s.get("sha1") in known:
                continue
            reg.append({"file": f, "sha1": s.get("sha1"),
                        "folder": s.get("folder"),
                        "imported_on": date.today().isoformat(),
                        "proposals": [i["id"] for i in added
                                      if i.get("source") == f]})
            if archive_root:
                m = archive(archive_root, spec, f)
                if m:
                    moved.append(m)
        lg.save()
    else:
        kept = [s.get("file", "") for s in (draft.get("sources") or [])
                if s.get("file") not in done_ok]

    return {"added": added, "duplicates": dup, "incomplete": bad,
            "moved": moved, "kept": kept}


def duplicate_of(lg, target: str, angle: str):
    """同じ対象 × 同じ切り口が既にあるか。無ければ None。"""
    if not target or not angle:
        return None
    t = _norm_target(target)
    for i in lg.items:
        if i.get("angle") != angle:
            continue
        if i.get("target") and _norm_target(i["target"]) == t:
            return i
    return None


def _norm_target(v: str) -> str:
    """対象URLの表記ゆれを吸収する。末尾スラッシュ、scheme、www、クエリ。"""
    v = (v or "").strip().lower()
    v = re.sub(r"^https?://", "", v)
    v = re.sub(r"^www\.", "", v)
    v = v.split("?", 1)[0].split("#", 1)[0]
    return v.rstrip("/") or "/"


def archive(root: str, spec: dict, rel: str) -> str:
    """原本を _archive/ へ退避する。以後は台帳だけを読む。"""
    base = os.path.join(root, spec.get("dir_name", "_proposals_inbox"))
    src = os.path.join(base, rel.replace("/", os.sep))
    if not os.path.exists(src):
        return ""
    dst_dir = os.path.join(base, spec.get("archive_dir", "_archive"),
                           os.path.dirname(rel.replace("/", os.sep)))
    os.makedirs(dst_dir, exist_ok=True)
    dst = os.path.join(dst_dir, os.path.basename(src))
    if os.path.exists(dst):
        stem, ext = os.path.splitext(os.path.basename(src))
        dst = os.path.join(dst_dir, f"{stem}_{date.today():%Y%m%d}{ext}")
    shutil.move(src, dst)
    return dst


# ---------------------------------------------------------------- 前段チェック
def warnings(lg, root: str, spec: dict, period: str) -> list[str]:
    """レポートを作る前に出す警告。

    「毎月必ず読む」を運用ルールにすると守られない。**機械が前段で言う。**
    """
    out = []
    done = {r.get("sha1") for r in (lg.data.get("imports") or [])}
    found, _skipped = scan(root, spec, done)
    if found:
        boxes = sorted({r["folder"] for r in found})
        out.append(
            f"未取り込みの原本が {len(found)} 件あります（{'、'.join(boxes)}）。"
            "先に `ledger.py --import` を実行してください。"
            "取り込まないまま提案を書くと、実施済みの案を出し直すことになります。")
    months = ((spec.get("staleness") or {}).get("warn_after_months")) or 3
    last = last_touched(lg)
    if lg.items and last and _months_between(last, period) >= months:
        out.append(
            f"台帳が {last} から更新されていません（{months} か月以上）。"
            "実施状況を確かめないまま提案を出すと、"
            "「過去の経緯を確認していない」という印象になります。"
            f"`ledger.py --status-sheet {period}` を打ち合わせに持っていってください。")
    if not lg.items:
        out.append(
            "台帳が空です。過去の提案を1件も把握していない状態です。"
            "`ledger.py --init-inbox` で取り込み口を作り、"
            "**まず自社の過去レポートから初期投入してください。**"
            "クライアントに書き出してもらう前提にすると、5つの箱が空のまま止まります。")
    return out


def last_touched(lg) -> str:
    """台帳が最後に動いた月。作成・実装・取り込みのどれでもよい。"""
    seen = []
    for i in lg.items:
        for k in ("created", "implemented_on", "imported_on"):
            v = i.get(k)
            if v:
                seen.append(str(v)[:7])
    for r in (lg.data.get("imports") or []):
        if r.get("imported_on"):
            seen.append(str(r["imported_on"])[:7])
    return max(seen) if seen else ""


def _months_between(a: str, b: str) -> int:
    try:
        ay, am = int(a[:4]), int(a[5:7])
        by, bm = int(b[:4]), int(b[5:7])
    except (ValueError, IndexError):
        return 0
    return (by - ay) * 12 + (bm - am)
