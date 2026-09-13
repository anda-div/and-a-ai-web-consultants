# -*- coding: utf-8 -*-
"""過去資料の取り込みが、通すべきものを通し、止めるべきものを止めることを確かめる。

    python tests/test_ledger_import.py

この取り込みは**台帳に書き込む**ため、間違いが静かに積もる。
特に危ないのは次の2つで、どちらも「動いているように見えて壊れている」型である。

    ・埋まらなかった原本を _archive/ へ退避してしまう
      → 台帳に入っていないのに入口から消え、記録もされるので二度と拾われない。
        情報が黙って消える。
    ・重複をタイトルで判定してしまう
      → 言い換えられた同じ提案が素通りし、来月クライアントへもう一度出る。

python-pptx / python-docx / openpyxl / pypdf が無い環境では、
その形式のテストだけを飛ばす。
"""
import io
import json
import os
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))

import ledger as L            # noqa: E402
import ledger_import as LI    # noqa: E402

SPEC = L.load_json(L.INBOX_FILE)


def has(module: str) -> bool:
    try:
        __import__(module)
        return True
    except ImportError:
        return False


def put(inbox, folder, name, text="メモ"):
    p = os.path.join(inbox, folder, name)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with io.open(p, "w", encoding="utf-8") as f:
        f.write(text)
    return p


def row(**kw):
    base = {"source": "", "title": "題", "target": "/a", "angle": "x",
            "status": "提案中"}
    base.update(kw)
    return base


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = self.tmp.name
        os.makedirs(os.path.join(self.root, "_ledger"))
        self.lg = L.Ledger(os.path.join(self.root, "_ledger", "proposals.json"))
        LI.init_inbox(self.root, SPEC, "README")
        self.inbox = os.path.join(self.root, "_proposals_inbox")

    def tearDown(self):
        self.tmp.cleanup()

    def apply(self, proposals, sources=None, **kw):
        draft = {"sources": sources or [], "proposals": proposals}
        draft.update(kw)
        return LI.apply_draft(self.lg, draft, SPEC, period="2026-09",
                              archive_root=self.root)

    def scan(self):
        done = {r.get("sha1") for r in (self.lg.data.get("imports") or [])}
        return LI.scan(self.root, SPEC, done)


class 仕様(Base):
    """フォルダの定義そのもの。ここが崩れると全部ずれる。"""

    def test_箱の名前が状態の申告になる(self):
        put(self.inbox, "20_実装済み", "a.md")
        put(self.inbox, "30_却下・結論が出ていない", "b.md")
        found, _ = self.scan()
        got = {r["folder"]: r["status"] for r in found}
        self.assertEqual(got["20_実装済み"], "実装済み")
        # 却下ではなく保留が既定。「見送った」は「二度と出さない」ではない
        self.assertEqual(got["30_却下・結論が出ていない"], "保留")

    def test_説明用のキーを箱として扱わない(self):
        """JSONの _comment は文字列。仕様として舐めると落ちる。"""
        self.assertNotIn("_comment", LI.folder_specs(SPEC))
        for v in LI.folder_specs(SPEC).values():
            self.assertIsInstance(v, dict)

    def test_READMEは取り込み対象にしない(self):
        found, _ = self.scan()
        self.assertEqual(found, [])


class 重複(Base):
    """タイトルではなく target × angle で見る。"""

    def test_言い換えられた同じ提案を止める(self):
        self.apply([row(title="料金を分かりやすく表示する",
                        target="https://example.test/price/", angle="price")])
        res = self.apply([row(title="料金・最低保証を事前提示し、依頼フローを図解する",
                              target="http://www.example.test/price?utm=x",
                              angle="price")])
        self.assertEqual(len(res["added"]), 0)
        self.assertEqual(len(res["duplicates"]), 1)

    def test_切り口が違えば別の提案として通す(self):
        self.apply([row(target="/price", angle="price")])
        res = self.apply([row(target="/price", angle="form_friction")])
        self.assertEqual(len(res["added"]), 1)

    def test_対象が違えば別の提案として通す(self):
        self.apply([row(target="/price", angle="price")])
        res = self.apply([row(target="/contact", angle="price")])
        self.assertEqual(len(res["added"]), 1)


class 必須項目(Base):
    """推測で埋めない。埋まらないものは人に聞く。"""

    def test_切り口が空なら入れない(self):
        res = self.apply([row(angle="")])
        self.assertEqual(res["added"], [])
        self.assertIn("angle", res["incomplete"][0][1])

    def test_対象が空なら入れない(self):
        res = self.apply([row(target="")])
        self.assertIn("target", res["incomplete"][0][1])

    def test_実装済みは実装日が要る(self):
        """実装日が無いと前後比較ができず、台帳の価値が半分になる。"""
        res = self.apply([row(status="実装済み")])
        self.assertIn("implemented_on", res["incomplete"][0][1])

    def test_実装日は概算でよい(self):
        res = self.apply([row(status="実装済み", implemented_on="2025-10")])
        self.assertEqual(len(res["added"]), 1)

    def test_保留は何待ちかと期日が要る(self):
        """「次期リニューアル時に」は却下ではなく保留＋期日。"""
        res = self.apply([row(status="保留")])
        self.assertEqual(sorted(res["incomplete"][0][1]),
                         ["blocked_by", "revisit_on"])
        ok = self.apply([row(status="保留", blocked_by="要件確定",
                             revisit_on="2026-12")])
        self.assertEqual(len(ok["added"]), 1)


class 退避(Base):
    """**埋まらなかった原本を退避してはいけない。** 情報が黙って消える。"""

    def src(self, name, folder="10_過去の定期レポート"):
        p = put(self.inbox, folder, name)
        return {"file": f"{folder}/{name}", "sha1": LI.sha1(p), "folder": folder}

    def test_台帳に入ったものは退避する(self):
        s = self.src("ok.md")
        res = self.apply([row(source=s["file"])], sources=[s])
        self.assertEqual(len(res["moved"]), 1)
        self.assertEqual(res["kept"], [])

    def test_埋まらなかったものは入口に残す(self):
        s = self.src("ng.md")
        res = self.apply([row(source=s["file"], angle="")], sources=[s])
        self.assertEqual(res["moved"], [])
        self.assertEqual(res["kept"], [s["file"]])
        self.assertTrue(os.path.exists(
            os.path.join(self.inbox, s["file"].replace("/", os.sep))))

    def test_入口に残したものは次回また出てくる(self):
        s = self.src("ng.md")
        self.apply([row(source=s["file"], angle="")], sources=[s])
        found, _ = self.scan()
        self.assertEqual([r["file"] for r in found], [s["file"]])

    def test_すでに台帳にあるものは退避してよい(self):
        """重複なら中身は失われていない。"""
        self.apply([row(target="/p", angle="a")])
        s = self.src("dup.md")
        res = self.apply([row(source=s["file"], target="/p", angle="a")],
                         sources=[s])
        self.assertEqual(len(res["moved"]), 1)

    def test_提案なしと書けば退避する(self):
        """読んだうえで提案が無かったもの。書かないと永久に出続ける。"""
        s = self.src("none.md")
        res = self.apply([], sources=[s], **{"提案なし": [s["file"]]})
        self.assertEqual(len(res["moved"]), 1)

    def test_同じ原本を二度読まない(self):
        s = self.src("once.md")
        self.apply([row(source=s["file"])], sources=[s])
        found, skipped = self.scan()
        self.assertEqual(found, [])
        self.assertEqual(len(self.lg.data["imports"]), 1)


class 前段チェック(Base):
    """「毎月必ず読む」は守られない。機械が前段で言う。"""

    def test_台帳が空なら警告する(self):
        w = LI.warnings(self.lg, self.root, SPEC, "2026-09")
        self.assertTrue(any("台帳が空" in x for x in w))

    def test_未取り込みの原本があれば警告する(self):
        put(self.inbox, "10_過去の定期レポート", "a.md")
        w = LI.warnings(self.lg, self.root, SPEC, "2026-09")
        self.assertTrue(any("未取り込み" in x for x in w))

    def test_放置された台帳を警告する(self):
        # 取り込みを経ずに古い提案だけがある状態。取り込んだ当日は
        # 「台帳を更新した」ことになるので、ここでは直接足す
        self.lg.add(period="2026-01", title="古い提案", target="/a", angle="x")
        w = LI.warnings(self.lg, self.root, SPEC, "2026-09")
        self.assertTrue(any("更新されていません" in x for x in w))

    def test_取り込んだ直後は放置とみなさない(self):
        """取り込みも台帳の更新である。"""
        self.apply([row(created="2026-01")])
        w = LI.warnings(self.lg, self.root, SPEC, "2026-09")
        self.assertFalse(any("更新されていません" in x for x in w))

    def test_片付いていれば黙る(self):
        self.apply([row(created="2026-09")])
        self.assertEqual(LI.warnings(self.lg, self.root, SPEC, "2026-09"), [])


class 確認シート(Base):
    """取り込んだ提案は工数が空。3つの費用帯だけで組むと1件も載らない。"""

    def test_工数の無い提案も確認シートに載る(self):
        self.apply([row(title="過去に出していた提案", effort="")])
        sheet = L.status_sheet(self.lg, "2026-09")
        self.assertIn("過去に出していた提案", sheet)


class 読み取り(Base):
    """6形式。読めないことを黙って0文字として扱わない。"""

    def test_テキストとMarkdownを読む(self):
        p = put(self.inbox, "00_未分類", "a.md", "対象 /price の料金表示")
        text, why = LI.extract(p)
        self.assertEqual(why, "")
        self.assertIn("/price", text)

    def test_cp932のテキストも読む(self):
        p = os.path.join(self.inbox, "00_未分類", "sjis.txt")
        with io.open(p, "w", encoding="cp932") as f:
            f.write("・問い合わせボタンの移動　2025年11月")
        text, why = LI.extract(p)
        self.assertEqual(why, "")
        self.assertIn("問い合わせボタン", text)

    @unittest.skipUnless(has("pptx"), "python-pptx が無い環境のため飛ばす")
    def test_PowerPointは本文もノートも読む(self):
        from pptx import Presentation
        from pptx.util import Cm
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_masters[0].slide_layouts[6])
        s.shapes.add_textbox(Cm(2), Cm(2), Cm(10), Cm(2)).text_frame.text = "本文の提案"
        s.notes_slide.notes_text_frame.text = "会議で保留になった"
        p = os.path.join(self.inbox, "00_未分類", "a.pptx")
        prs.save(p)
        text, why = LI.extract(p)
        self.assertEqual(why, "")
        self.assertIn("本文の提案", text)
        self.assertIn("会議で保留になった", text)   # 経緯はノートに書かれる

    @unittest.skipUnless(has("docx"), "python-docx が無い環境のため飛ばす")
    def test_Wordは表も読む(self):
        import docx
        d = docx.Document()
        d.add_paragraph("段落の提案")
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "対象"
        t.rows[0].cells[1].text = "/faq"
        p = os.path.join(self.inbox, "00_未分類", "a.docx")
        d.save(p)
        text, _ = LI.extract(p)
        self.assertIn("段落の提案", text)
        self.assertIn("/faq", text)

    @unittest.skipUnless(has("openpyxl"), "openpyxl が無い環境のため飛ばす")
    def test_Excelを読む(self):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(["会員ページに導線", "/mypage"])
        p = os.path.join(self.inbox, "00_未分類", "a.xlsx")
        wb.save(p)
        text, _ = LI.extract(p)
        self.assertIn("/mypage", text)

    def test_読めない形式は理由を残す(self):
        p = put(self.inbox, "00_未分類", "a.zzz", "x")
        text, why = LI.extract(p)
        self.assertEqual(text, "")
        self.assertTrue(why)

    def test_壊れたファイルも理由を残して落ちない(self):
        p = put(self.inbox, "00_未分類", "broken.xlsx", "これはExcelではない")
        text, why = LI.extract(p)
        self.assertEqual(text, "")
        self.assertTrue(why)


class 手がかり(Base):
    """確定はしない。候補を並べてAIに選ばせる。"""

    def test_URLと日付を拾う(self):
        h = LI.hints("対象は https://example.test/price です。2025年10月に実施。", None)
        self.assertIn("https://example.test/price", h["targets"])
        self.assertIn("2025-10", h["dates"])

    def test_切り口はカタログと突き合わせる(self):
        cat = {"angles": [{"key": "price_transparency", "name": "料金の透明性"}]}
        h = LI.hints("切り口は price_transparency", cat)
        self.assertEqual(h["angles"], ["price_transparency"])

    def test_ありえない月は拾わない(self):
        self.assertEqual(LI.hints("2025年13月", None)["dates"], [])


if __name__ == "__main__":
    unittest.main(verbosity=2)
