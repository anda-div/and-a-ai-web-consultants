# -*- coding: utf-8 -*-
"""実施状況シート（Excel）── 出す・お客様が直す・取り込む、を通しで確かめる。

    python tests/test_status_sheet.py

危ないのは2つ。

    ・お客様が書いていないことを、台帳に書き込む（実装日を今日で埋める、など）
    ・お客様の取り違えで、済んだ提案が未着手に戻る

どちらも台帳の中では辻褄が合ってしまい、翌月の前後比較が黙って狂う。
"""
import contextlib
import io
import os
import subprocess
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))

from openpyxl import load_workbook  # noqa: E402

import ledger as L  # noqa: E402
import ledger_import  # noqa: E402
import status_sheet_xlsx as S  # noqa: E402


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = self.tmp.name
        self.lg = L.Ledger(os.path.join(self.root, "_ledger", "proposals.json"))
        self.path = os.path.join(self.root, "_input", "実施状況", "シート.xlsx")

    def tearDown(self):
        self.tmp.cleanup()

    def add(self, **kw):
        base = dict(period="2026-08", title="題", target="/form", angle="cta")
        base.update(kw)
        return self.lg.add(**base)

    def export(self, period="2026-10"):
        return S.export(self.lg, period, self.path, L.unconfirmed_import)

    def edit(self, pid, **cols):
        """お客様が黄色の列を直したことにする。"""
        wb = load_workbook(self.path)
        ws = wb[S.SHEET_NAME]
        head = {c.value: c.column for c in ws[S.HEADER_ROW]}
        for row in ws.iter_rows(min_row=S.HEADER_ROW + 1):
            if row[0].value == pid:
                for k, v in cols.items():
                    ws.cell(row=row[0].row, column=head[k], value=v)
        wb.save(self.path)

    def diff(self, period="2026-10"):
        _p, rows = S.read(self.path)
        return S.diff(self.lg, rows, period=period)

    def ids_in_sheet(self):
        return [r["提案ID"] for r in S.read(self.path)[1]]


class 出力(Base):
    def test_決着したものは載せず_いまの状態を記入して出す(self):
        a = self.add(title="未着手のもの")
        b = self.add(title="却下のもの", status="却下")
        c = self.add(title="対応中のもの", status="実装待ち")
        self.export()
        self.assertEqual(sorted(self.ids_in_sheet()), sorted([a["id"], c["id"]]))
        rows = {r["提案ID"]: r for r in S.read(self.path)[1]}
        self.assertEqual(rows[a["id"]]["状態"], "未着手")
        self.assertEqual(rows[c["id"]]["状態"], "対応中")
        self.assertNotIn(b["id"], rows)

    def test_状態はプルダウンで_突き合わせ用の列は隠す(self):
        self.add()
        self.export()
        ws = load_workbook(self.path)[S.SHEET_NAME]
        lists = [dv.formula1 for dv in ws.data_validations.dataValidation]
        self.assertTrue(any("実装済み（提案前から）" in f for f in lists))
        col = S.HEADERS.index("出力時の状態") + 1
        self.assertTrue(ws.column_dimensions[ws.cell(row=1, column=col).column_letter].hidden)

    def test_実装日の分からない実装済みを先頭に出して聞き直す(self):
        self.add(title="ふつう")
        x = self.add(title="日付不明", status="実装済み")
        self.lg.set_status(x["id"], "実装済み", date_unknown=True)
        self.export()
        self.assertEqual(self.ids_in_sheet()[0], x["id"])


class 取り込み(Base):
    def test_直していないシートは何も変えない(self):
        self.add()
        self.add(status="実装待ち")
        self.export()
        res = self.diff()
        self.assertEqual(res["changes"], [])

    def test_実装済みと実装日を入れる_経緯も残る(self):
        a = self.add()
        self.export()
        self.edit(a["id"], 状態="実装済み", 実装日="2026/9/15")
        res = self.diff()
        S.apply(self.lg, res["changes"], source="実施状況シート 2026-10")
        self.assertEqual(a["status"], "実装済み")
        self.assertEqual(a["implemented_on"], "2026-09-15")
        self.assertEqual(a["status_log"][-1]["source"], "実施状況シート 2026-10")

    def test_実装日が無ければ今日で埋めない(self):
        a = self.add()
        self.export()
        self.edit(a["id"], 状態="実装済み")
        S.apply(self.lg, self.diff()["changes"], source="t")
        self.assertEqual(a["status"], "実装済み")
        self.assertIsNone(a["implemented_on"])
        self.assertEqual(self.lg.verify_due("2026-11"), [])

    def test_おおよその実装日は月まで受ける(self):
        self.assertEqual(S.parse_date("2026年9月中旬"), ("2026-09", ""))
        self.assertEqual(S.parse_date("2026-09"), ("2026-09", ""))
        self.assertEqual(S.parse_date("２０２６／９／３"), ("2026-09-03", ""))
        self.assertEqual(S.parse_date("わからない"), ("", "わからない"))

    def test_読めない実装日は推測せず確認待ちにする(self):
        a = self.add()
        self.export()
        self.edit(a["id"], 状態="実装済み", 実装日="夏ごろ")
        res = self.diff()
        self.assertEqual(res["unreadable"], [(a["id"], "夏ごろ")])
        self.assertTrue(res["changes"][0]["date_unknown"])

    def test_提案前から実装済みは前後比較せず_次のシートにも出さない(self):
        a = self.add()
        self.export()
        self.edit(a["id"], 状態="実装済み（提案前から）")
        S.apply(self.lg, self.diff()["changes"], source="t")
        self.assertTrue(a["implemented_before_proposal"])
        self.assertIsNone(a["implemented_on"])
        self.assertIn("提案前から実装済み", a["status_note"])
        self.export()
        self.assertNotIn(a["id"], self.ids_in_sheet())

    def test_済んだものを未着手に戻す記入は反映しない(self):
        a = self.add(status="実装済み", implemented_on="2026-09-01")
        self.export()
        self.edit(a["id"], 状態="未着手")
        res = self.diff()
        self.assertEqual(res["changes"], [])
        self.assertEqual([c["id"] for c in res["conflicts"]], [a["id"]])
        self.assertEqual(a["status"], "実装済み")

    def test_却下だったものも_お客様が実装済みと書けば正とする(self):
        a = self.add(status="提案中")
        self.export()
        a["status"] = "却下"          # シートを出したあとに台帳だけ動いた
        self.edit(a["id"], 状態="実装済み", 実装日="2026-10-01")
        res = self.diff()
        self.assertEqual(res["changes"][0]["to"], "実装済み")
        self.assertTrue(any("却下" in r for r in res["changes"][0]["remarks"]))

    def test_やらないは却下にし_理由が無ければ知らせる(self):
        a = self.add()
        self.export()
        self.edit(a["id"], 状態="やらない")
        res = self.diff()
        self.assertEqual(res["changes"][0]["to"], "却下")
        self.assertTrue(any("理由" in r for r in res["changes"][0]["remarks"]))

    def test_保留は補足を何待ちとし_時期は既定で埋める(self):
        a = self.add()
        self.export()
        self.edit(a["id"], 状態="保留", 補足="リニューアル待ち")
        S.apply(self.lg, self.diff(period="2026-10")["changes"], source="t")
        self.assertEqual(a["status"], "保留")
        self.assertEqual(a["blocked_by"], "リニューアル待ち")
        self.assertEqual(a["revisit_on"], "2027-01")

    def test_台帳に無いIDは知らせる(self):
        a = self.add()
        self.export()
        self.edit(a["id"], 提案ID="P-1999-01-001")
        self.assertEqual(self.diff()["unknown"], ["P-1999-01-001"])


class 実装日を作り話にしない(Base):
    def test_日付を省いて実装済みにすると止まる(self):
        a = self.add()
        with self.assertRaises(ValueError):
            self.lg.set_status(a["id"], "実装済み")
        self.assertEqual(a["status"], "提案中")      # 止まったときは何も変えない
        self.assertNotIn("status_log", a)

    def test_不明なら空のまま_日付があればその日(self):
        a = self.add()
        b = self.add()
        self.lg.set_status(a["id"], "実装済み", date_unknown=True)
        self.lg.set_status(b["id"], "実装済み", on="2026-09")
        self.assertIsNone(a["implemented_on"])
        self.assertEqual(b["implemented_on"], "2026-09")

    def test_コマンドでも_日付を省くと止まり台帳は変わらない(self):
        a = self.add()
        self.lg.save()
        r = subprocess.run(
            [sys.executable, os.path.join(ROOT, "scripts", "ledger.py"),
             "--path", self.lg.path, "--set", a["id"], "--status", "実装済み"],
            capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(r.returncode, 1)
        self.assertIn("--date 不明", r.stdout)
        self.assertEqual(L.Ledger(self.lg.path).get(a["id"])["status"], "提案中")


class こちらの理解(Base):
    def test_状態の補足があれば_それを出す(self):
        a = self.add(status_note="中身を変えて残すと伺っています。改訂版をお持ちします")
        b = self.add()
        self.export()
        rows = {r["提案ID"]: r for r in S.read(self.path)[1]}
        self.assertEqual(rows[a["id"]]["こちらの理解"], a["status_note"])
        self.assertEqual(rows[b["id"]]["こちらの理解"], "まだご返答をいただいていません")
        self.assertIn(a["status_note"], L.status_sheet(self.lg, "2026-10"))


class 実装済みの可能性あり(Base):
    def test_見たものを書かないと付けられない(self):
        a = self.add()
        with self.assertRaises(ValueError):
            self.lg.mark_maybe_done(a["id"], "")

    def test_状態は変えず_シートで尋ね_確かめたら外れる(self):
        a = self.add()
        self.lg.mark_maybe_done(a["id"], "2026-10 に撮り直したら直っていた")
        self.assertEqual(a["status"], "提案中")
        self.export()
        row = S.read(self.path)[1][0]
        self.assertIn("実装済みの可能性", row["こちらの理解"])
        self.lg.set_status(a["id"], "実装済み", on="2026-09-20")
        self.assertNotIn("maybe_done", a)


class レポート前の確認(Base):
    def test_戻ったシートを取り込むまで警告し_取り込めば消える(self):
        a = self.add(period="2026-10")
        self.export()
        warns = ledger_import.warnings(self.lg, self.root, {}, "2026-10")
        self.assertTrue(any("実施状況シート" in w for w in warns))
        self.lg.data.setdefault("status_sheets", []).append(
            {"sha1": S.sha1(self.path), "imported_on": "2026-10-18"})
        warns = ledger_import.warnings(self.lg, self.root, {}, "2026-10")
        self.assertFalse(any("実施状況シート" in w for w in warns))
        self.assertTrue(a)

    def test_シートが戻っていなければ知らせるが止めない(self):
        self.lg.data["status_sheet_exports"] = [
            {"period": "2026-10", "exported_on": "2026-10-01"}]
        self.assertIn("まだ戻っていません", ledger_import.sheet_notice(self.lg))
        self.lg.data["status_sheets"] = [{"imported_on": "2026-10-05"}]
        self.assertEqual(ledger_import.sheet_notice(self.lg), "")

    def test_同じシートを2度は入れない(self):
        a = self.add()
        self.export()
        self.edit(a["id"], 状態="対応中")
        with contextlib.redirect_stdout(io.StringIO()):
            L.import_sheet(self.lg, self.path, apply=True)
            n = len(a["status_log"])
            L.import_sheet(self.lg, self.path, apply=True)
        self.assertEqual(len(a["status_log"]), n)


class 在庫ページ(Base):
    def test_件数と保留の時期を出す(self):
        self.add()
        h = self.add(title="保留のもの")
        h.update(status="保留", blocked_by="予算", revisit_on="2027-01")
        text = L.inventory(self.lg, "2026-10", os.path.dirname(self.lg.path))
        self.assertIn("| 未着手 | 1 |", text)
        self.assertIn("2027-01", text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
